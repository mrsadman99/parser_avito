from pathlib import Path
from threading import Lock
from datetime import datetime
import re

from openpyxl import Workbook, load_workbook
from loguru import logger
from tzlocal import get_localzone

from parser.export.base import ResultStorage
from models import Item

SCAN_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")

class ExcelStorage(ResultStorage):
    """
    Сохранение результатов парсинга в XLSX
    """
    name = "excel"
    headers = [
        "Название",
        "Цена",
        "URL",
        "Описание",
        "Дата публикации",
        "Дата сканирования",
        "Продавец",
        "Адрес",
        "Адрес пользователя",
        "Координаты",
        "Изображения",
        "Поднято",
        "Телефон",
        "Рейтинг продавца",
        "Кол-во оценок",
    ]

    def __init__(self, file_path: Path):
        self.file_path = file_path

        # создаём директорию
        self.file_path.parent.mkdir(parents=True, exist_ok=True)

        # lock для потокобезопасной записи
        self._lock = Lock()

        # создаём файл, если его нет
        if not self.file_path.exists():
            self._create_file()
        else:
            self._ensure_schema()

    def _create_file(self) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet.append(self.headers)
        workbook.save(self.file_path)

    def _ensure_schema(self) -> None:
        """Приводит существующий файл к актуальному набору колонок.

        - добавляет недостающую колонку «Дата сканирования» в старых файлах;
        - удаляет лишние колонки справа (например, AI-колонки), чтобы шапка
          и строки были однородными.
        """
        try:
            workbook = load_workbook(self.file_path)
            sheet = workbook.active

            header = [cell.value for cell in sheet[1]][: len(self.headers)]
            extra_cols = sheet.max_column != len(self.headers)
            if header == self.headers and not extra_cols:
                return

            max_col = len(self.headers)

            # 1) нормализуем строки данных
            for row_idx in range(2, sheet.max_row + 1):
                values = [cell.value for cell in sheet[row_idx]]
                if not any(v is not None for v in values):
                    continue
                # в старых файлах нет колонки сканирования — вставляем её
                scan_cell = values[5] if len(values) > 5 else None
                if not (isinstance(scan_cell, str) and SCAN_DATE_RE.match(scan_cell)):
                    values = values[:5] + [None] + values[5:]
                values = values[:max_col]
                values += [None] * (max_col - len(values))
                for col_idx, value in enumerate(values, start=1):
                    # присваиваем через .value: cell(..., value=None) не очищает ячейку
                    sheet.cell(row=row_idx, column=col_idx).value = value

            # 2) удаляем лишние колонки справа
            if sheet.max_column > max_col:
                sheet.delete_cols(max_col + 1, sheet.max_column - max_col)

            # 3) переписываем шапку
            for idx, title in enumerate(self.headers, start=1):
                sheet.cell(row=1, column=idx).value = title

            workbook.save(self.file_path)
            logger.info("Excel: шапка приведена к актуальным колонкам")
        except Exception as err:
            logger.warning(f"Excel: не удалось проверить шапку файла: {err}")

    @staticmethod
    def _get_ad_time(ad: Item):
        return (
            datetime
            .fromtimestamp(ad.sortTimeStamp / 1000, tz=get_localzone())
            .replace(tzinfo=None)
        )

    @staticmethod
    def _get_scanned_at(ad: Item) -> str:
        """Дата сканирования: из UTC (scanned_at) в локальное время пользователя."""
        scanned_at = getattr(ad, "scanned_at", None)
        if not scanned_at:
            return datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        try:
            dt = datetime.fromisoformat(scanned_at)
            return dt.astimezone().strftime("%Y-%m-%d %H:%M:%S")
        except (ValueError, TypeError):
            return str(scanned_at)

    @staticmethod
    def _get_item_coords(ad: Item) -> str:
        if ad.coords and "lat" in ad.coords and "lng" in ad.coords:
            return f"{ad.coords['lat']};{ad.coords['lng']}"
        return ""

    @staticmethod
    def _get_item_address_user(ad: Item) -> str:
        if ad.coords and "address_user" in ad.coords:
            return ad.coords["address_user"]
        return ""

    @staticmethod
    def _get_largest_image_url(img) -> str:
        try:
            best_key = max(
                img.root.keys(),
                key=lambda k: int(k.split("x")[0]) * int(k.split("x")[1])
            )
            return str(img.root[best_key])
        except Exception as err:
            logger.error(f"При определении лучшего изображения ошибка: {err}")
            return ""

    @staticmethod
    def excel_safe(value):
        # Formula Injection fix
        if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
            return "'" + value
        return value

    def save(self, ads: list[Item]) -> None:
        if not ads:
            return

        with self._lock:
            workbook = load_workbook(self.file_path)
            sheet = workbook.active

            for ad in ads:
                images_urls = [
                    self._get_largest_image_url(img)
                    for img in ad.images
                ]

                row = [
                    self.excel_safe(ad.title),
                    ad.priceDetailed.value,
                    self.excel_safe(f"https://www.avito.ru/{ad.urlPath}"),
                    self.excel_safe(ad.description),
                    self._get_ad_time(ad),
                    self.excel_safe(self._get_scanned_at(ad)),
                    self.excel_safe(ad.sellerId or ""),
                    self.excel_safe(ad.location.name if ad.location else ""),
                    self.excel_safe(self._get_item_address_user(ad)),
                    self.excel_safe(self._get_item_coords(ad)),
                    self.excel_safe(";".join(images_urls)),
                    "Да" if ad.isPromotion else "Нет",
                    self.excel_safe(ad.phone or ""),
                    ad.seller_rating if ad.seller_rating is not None else "",
                    ad.seller_reviews if ad.seller_reviews is not None else "",
                ]

                sheet.append(row)

            workbook.save(self.file_path)



